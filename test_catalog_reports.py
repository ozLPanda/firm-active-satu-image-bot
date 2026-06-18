from __future__ import annotations

from pathlib import Path
import os
import tempfile
import unittest

from openpyxl import Workbook

import app as bot_app
from bot import DeferredUploadItem, ExistingImageItem, ProductCandidate, SatuImageBot
from catalog import CatalogItem, ImageItem, image_code_from_path, is_supported_image, load_catalog, names_match, normalize_availability, processing_items, scan_images
from reports import ReportRecord, ReportWriter


class CatalogReportsTests(unittest.TestCase):
    def test_image_code_from_path(self) -> None:
        self.assertEqual(image_code_from_path(Path("00000001200.jpeg")), "00000001200")

    def test_supported_image_extensions(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            image_path = folder / "photo.JPG"
            image_path.write_text("", encoding="utf-8")
            self.assertTrue(is_supported_image(image_path))
            self.assertFalse(is_supported_image(folder / "photo.bmp"))

    def test_names_match_with_minor_difference(self) -> None:
        self.assertTrue(names_match("Котел газовый настенный 24 кВт", "Котел настенный газовый 24 кВт"))

    def test_load_catalog_detects_columns(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "catalog.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Наименование", "Цена", "Наличие"])
            sheet.append(["00000002416", "Котел", 240.0, "В наличии"])
            workbook.save(path)

            catalog = load_catalog(path)

            self.assertEqual(catalog["00000002416"].name, "Котел")
            self.assertEqual(catalog["00000002416"].price, "240")
            self.assertEqual(catalog["00000002416"].availability, "В наличии")

    def test_scan_images_attaches_catalog_price(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            image_path = folder / "00000002416.jpg"
            image_path.write_text("", encoding="utf-8")
            path = folder / "catalog.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Наименование", "Цена", "Наличие"])
            sheet.append(["00000002416", "Котел", "1 250", "Нет в наличии"])
            workbook.save(path)

            catalog = load_catalog(path)
            images = scan_images(folder, catalog)

            self.assertEqual(images[0].price, "1250")
            self.assertEqual(images[0].availability, "Нет в наличии")

    def test_processing_items_includes_catalog_rows_without_images(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            path = folder / "catalog.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Наименование", "Цена", "Наличие"])
            sheet.append(["00000002416", "Котел", "1 250", "В наличии"])
            workbook.save(path)

            catalog = load_catalog(path)
            items = processing_items(folder, catalog)

            self.assertEqual(items[0].code, "00000002416")
            self.assertIsNone(items[0].path)
            self.assertEqual(items[0].price, "1250")

    def test_processing_items_matches_images_without_leading_zeroes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            image_path = folder / "2416.jpg"
            image_path.write_text("", encoding="utf-8")
            catalog = {
                "00000002416": CatalogItem(
                    code="00000002416",
                    name="Pump",
                    price="1250",
                    availability="В наличии",
                )
            }

            items = processing_items(folder, catalog)

            self.assertEqual(items[0].code, "00000002416")
            self.assertEqual(items[0].path, image_path)

    def test_catalog_infers_unavailable_from_name_marker(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "catalog.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Наименование", "Цена"])
            sheet.append(["00000002416", "НЕТ В НАЛИЧИИ Котел", 240])
            workbook.save(path)

            catalog = load_catalog(path)

            self.assertEqual(catalog["00000002416"].name, "Котел")
            self.assertEqual(catalog["00000002416"].availability, "Нет в наличии")

    def test_catalog_infers_preorder_from_name_marker(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "catalog.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Артикул", "Наименование", "Цена"])
            sheet.append(["00000002416", "ПОД ЗАКАЗ Котел", 240])
            workbook.save(path)

            catalog = load_catalog(path)

            self.assertEqual(catalog["00000002416"].name, "Котел")
            self.assertEqual(catalog["00000002416"].availability, "Под заказ")

    def test_normalize_availability_from_stock_values(self) -> None:
        self.assertEqual(normalize_availability(3), "В наличии")
        self.assertEqual(normalize_availability(0), "Нет в наличии")
        self.assertEqual(normalize_availability("под заказ"), "Под заказ")

    def test_report_writer_saves_csv_and_json(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            writer = ReportWriter(temp_dir)
            writer.add(ReportRecord(code="0001", status="uploaded", image_path="0001.jpg"))

            self.assertTrue(writer.csv_path.exists())
            self.assertTrue(writer.json_path.exists())
            self.assertIn("uploaded", writer.csv_path.read_text(encoding="utf-8-sig"))
            self.assertIn("uploaded", writer.json_path.read_text(encoding="utf-8"))

    def test_report_writer_keeps_ten_latest_report_sets(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            for index in range(12):
                csv_path = folder / f"upload-report-20260101-0000{index:02d}.csv"
                json_path = folder / f"upload-report-20260101-0000{index:02d}.json"
                csv_path.write_text("", encoding="utf-8")
                json_path.write_text("", encoding="utf-8")
                timestamp = 1_700_000_000 + index
                os.utime(csv_path, (timestamp, timestamp))
                os.utime(json_path, (timestamp, timestamp))

            writer = ReportWriter(folder)
            writer.add(ReportRecord(code="0001", status="uploaded"))

            report_stems = {
                path.stem
                for path in folder.glob("upload-report-*.*")
                if path.suffix in {".csv", ".json"}
            }
            self.assertLessEqual(len(report_stems), 10)

    def test_report_writer_keeps_ten_latest_log_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            for index in range(12):
                path = folder / f"bot-{index:02d}.log"
                path.write_text("", encoding="utf-8")
                timestamp = 1_700_000_000 + index
                os.utime(path, (timestamp, timestamp))

            writer = ReportWriter(folder)
            writer.add(ReportRecord(code="0001", status="uploaded"))

            self.assertLessEqual(len(list(folder.glob("*.log"))), 10)

    def test_manual_queues_can_be_saved_and_trimmed(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            original_app_data_dir = bot_app.APP_DATA_DIR
            original_existing_queue_file = bot_app.EXISTING_QUEUE_FILE
            original_deferred_queue_file = bot_app.DEFERRED_QUEUE_FILE
            try:
                bot_app.APP_DATA_DIR = folder
                bot_app.EXISTING_QUEUE_FILE = folder / "existing.json"
                bot_app.DEFERRED_QUEUE_FILE = folder / "deferred.json"
                existing = ExistingImageItem(
                    code="0001",
                    image_path=folder / "0001.jpg",
                    product_url="https://example.test/product/1",
                    expected_name="Pump",
                    actual_name="Pump",
                    price="100",
                    availability="В наличии",
                )
                deferred = DeferredUploadItem(
                    code="0002",
                    image_path=None,
                    product_url="https://example.test/product/2",
                    expected_name="Valve",
                    actual_name="Valve",
                    message="save failed",
                    price="200",
                    availability="Нет в наличии",
                    image_upload_required=False,
                )

                bot_app.save_existing_queue([existing])
                bot_app.save_deferred_queue([deferred])
                self.assertEqual(len(bot_app.load_existing_queue()), 1)
                self.assertEqual(len(bot_app.load_deferred_queue()), 1)

                bot_app.remove_existing_queue_items([existing])
                bot_app.remove_deferred_queue_items([deferred])
                self.assertEqual(bot_app.load_existing_queue(), [])
                self.assertEqual(bot_app.load_deferred_queue(), [])
            finally:
                bot_app.APP_DATA_DIR = original_app_data_dir
                bot_app.EXISTING_QUEUE_FILE = original_existing_queue_file
                bot_app.DEFERRED_QUEUE_FILE = original_deferred_queue_file

    def test_saved_paths_are_loaded_only_when_targets_exist(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            image_folder = folder / "images"
            image_folder.mkdir()
            excel_path = folder / "catalog.xlsx"
            excel_path.write_text("", encoding="utf-8")
            original_app_data_dir = bot_app.APP_DATA_DIR
            original_user_settings_file = bot_app.USER_SETTINGS_FILE
            try:
                bot_app.APP_DATA_DIR = folder
                bot_app.USER_SETTINGS_FILE = folder / "user-settings.json"

                bot_app.save_saved_paths(image_folder=image_folder, excel_path=excel_path)
                self.assertEqual(
                    bot_app.load_saved_paths(),
                    {"image_folder": str(image_folder), "excel_path": str(excel_path)},
                )

                excel_path.unlink()
                self.assertEqual(
                    bot_app.load_saved_paths(),
                    {"image_folder": str(image_folder), "excel_path": ""},
                )
            finally:
                bot_app.APP_DATA_DIR = original_app_data_dir
                bot_app.USER_SETTINGS_FILE = original_user_settings_file

    def test_catalog_cache_entries_update_only_processed_codes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            original_app_data_dir = bot_app.APP_DATA_DIR
            original_catalog_cache_file = bot_app.CATALOG_CACHE_FILE
            try:
                bot_app.APP_DATA_DIR = folder
                bot_app.CATALOG_CACHE_FILE = folder / "catalog-cache.json"
                old_catalog = {
                    "0001": CatalogItem(code="0001", name="Old pump", price="100", availability="В наличии"),
                    "0002": CatalogItem(code="0002", name="Old valve", price="200", availability="В наличии"),
                }
                new_catalog = {
                    "0001": CatalogItem(code="0001", name="New pump", price="150", availability="Нет в наличии"),
                    "0002": CatalogItem(code="0002", name="New valve", price="250", availability="Нет в наличии"),
                }
                bot_app.save_catalog_cache(folder / "old.xlsx", old_catalog)

                bot_app.save_catalog_cache_entries(folder / "new.xlsx", new_catalog, {"0001"})
                cache = bot_app.load_catalog_cache()

                self.assertEqual(cache["products"]["0001"]["name"], "New pump")
                self.assertEqual(cache["products"]["0001"]["price"], "150")
                self.assertEqual(cache["products"]["0002"]["name"], "Old valve")
                self.assertEqual(cache["products"]["0002"]["price"], "200")
            finally:
                bot_app.APP_DATA_DIR = original_app_data_dir
                bot_app.CATALOG_CACHE_FILE = original_catalog_cache_file

    def test_successful_catalog_cache_codes_excludes_failed_records(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            writer = ReportWriter(temp_dir)
            writer.add(ReportRecord(code="0001", status="uploaded"))
            writer.add(ReportRecord(code="0002", status="not_found"))
            writer.add(ReportRecord(code="0003", status="deferred_upload_error"))
            existing = [
                ExistingImageItem(
                    code="0004",
                    image_path=Path(temp_dir) / "0004.jpg",
                    product_url="https://example.test/product/4",
                    expected_name="Pump",
                    actual_name="Pump",
                )
            ]

            self.assertEqual(bot_app.successful_catalog_cache_codes(writer, existing), {"0001", "0004"})

    def test_candidate_is_actual_when_list_row_matches(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot = SatuImageBot(temp_dir, ReportWriter(temp_dir))
            image = ImageItem(
                code="00000002416",
                path=Path(temp_dir) / "00000002416.jpg",
                catalog_name="Pump",
                price="1250",
                availability="В наличии",
            )
            candidate = ProductCandidate(
                url="https://my.satu.kz/cms/product/edit/1",
                text="Pump",
                code="00000002416",
                price="1 250 ₸",
                availability="В наличии",
                has_image=True,
            )

            self.assertTrue(bot._candidate_is_actual(image, candidate))

    def test_code_fields_prioritize_sku_input(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot = SatuImageBot(temp_dir, ReportWriter(temp_dir))

            self.assertEqual(bot.selectors.code_fields[0], "input[data-qaid='sku_input']")

    def test_codes_match_ignores_leading_zeroes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot = SatuImageBot(temp_dir, ReportWriter(temp_dir))

            self.assertTrue(bot._codes_match("2146", "00000002146"))

    def test_direct_link_lookup_matches_card_code_without_leading_zeroes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot = SatuImageBot(temp_dir, ReportWriter(temp_dir))
            image = ImageItem(
                code="00000002416",
                path=Path(temp_dir) / "00000002416.jpg",
                catalog_name="Pump",
                price="1250",
                availability="В наличии",
            )

            found = bot._direct_image_item_for_code("2416", {"00000002416": image})

            self.assertIs(found, image)

    def test_candidate_needs_card_when_image_missing_in_list_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot = SatuImageBot(temp_dir, ReportWriter(temp_dir))
            image = ImageItem(
                code="00000002416",
                path=Path(temp_dir) / "00000002416.jpg",
                catalog_name="Pump",
                price="1250",
                availability="В наличии",
            )
            candidate = ProductCandidate(
                url="https://my.satu.kz/cms/product/edit/1",
                text="Pump",
                code="00000002416",
                price="1 250 ₸",
                availability="В наличии",
                has_image=False,
            )

            self.assertFalse(bot._candidate_is_actual(image, candidate))

    def test_candidate_existing_image_can_be_queued_from_list_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot = SatuImageBot(temp_dir, ReportWriter(temp_dir))
            image = ImageItem(
                code="00000002416",
                path=Path(temp_dir) / "00000002416.jpg",
                catalog_name="Pump",
                price="1250",
                availability="В наличии",
            )
            candidate = ProductCandidate(
                url="https://my.satu.kz/cms/product/edit/1",
                text="Pump",
                code="00000002416",
                price="1 250 ₸",
                availability="В наличии",
                has_image=True,
            )

            self.assertTrue(bot._candidate_has_existing_image_for_replacement(image, candidate))

    def test_candidate_existing_image_needs_card_when_fields_differ(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            bot = SatuImageBot(temp_dir, ReportWriter(temp_dir))
            image = ImageItem(
                code="00000002416",
                path=Path(temp_dir) / "00000002416.jpg",
                catalog_name="Pump",
                price="1250",
                availability="В наличии",
            )
            candidate = ProductCandidate(
                url="https://my.satu.kz/cms/product/edit/1",
                text="Pump",
                code="00000002416",
                price="999 ₸",
                availability="В наличии",
                has_image=True,
            )

            self.assertFalse(bot._candidate_has_existing_image_for_replacement(image, candidate))


if __name__ == "__main__":
    unittest.main()
