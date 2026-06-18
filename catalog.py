from __future__ import annotations

from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
import re
from typing import Iterable

from openpyxl import load_workbook


SUPPORTED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}

CODE_HEADERS = {
    "код",
    "код товара",
    "код_товара",
    "артикул",
    "article",
    "sku",
}

NAME_HEADERS = {
    "название",
    "название позиции",
    "название_позиции",
    "наименование",
    "товар",
    "name",
    "product name",
}

PRICE_HEADERS = {
    "цена",
    "розничная цена",
    "цена розничная",
    "price",
}

WHOLESALE_PRICE_HEADERS = {
    "оптовая цена",
    "оптовая_цена",
}

SKO_PRICE_HEADERS = {
    "ско цена",
    "ско_цена",
}

NAME_RU_HEADERS = {
    "наименование ru",
    "наименование_ru",
    "name ru",
    "name_ru",
}

NAME_KZ_HEADERS = {
    "наименование kz",
    "наименование_kz",
    "наименование kk",
    "наименование_kk",
    "name kz",
    "name_kz",
    "name kk",
    "name_kk",
}

AVAILABILITY_HEADERS = {
    "наличие",
    "остаток",
    "статус",
    "availability",
    "stock",
}

NAME_AVAILABILITY_PATTERNS = (
    (
        re.compile(r"(?i)(?:\s*[\(\[\{,;:\-–—]\s*)?\bнет\s+в\s+наличии\b(?:\s*[\)\]\},;:\-–—]\s*)?"),
        "Нет в наличии",
    ),
    (
        re.compile(r"(?i)(?:\s*[\(\[\{,;:\-–—]\s*)?\bпод\s+заказ\b(?:\s*[\)\]\},;:\-–—]\s*)?"),
        "Под заказ",
    ),
)


@dataclass(frozen=True)
class CatalogItem:
    code: str
    name: str
    price: str | None = None
    wholesale_price: str | None = None
    sko_price: str | None = None
    availability: str | None = None


@dataclass(frozen=True)
class ImageItem:
    code: str
    path: Path | None
    catalog_name: str | None = None
    price: str | None = None
    wholesale_price: str | None = None
    sko_price: str | None = None
    name_ru: str | None = None
    name_kz: str | None = None
    availability: str | None = None


@dataclass(frozen=True)
class NameCatalogItem:
    code: str
    name_ru: str | None = None
    name_kz: str | None = None


def normalize_code(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def normalize_price(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"[\s\u00a0]+", "", text).replace(",", ".")
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def name_without_availability_marker(value: object) -> tuple[str, str | None]:
    text = "" if value is None else str(value).strip()
    text = re.sub(r"\s+", " ", text)
    detected_availability = None
    cleaned = text
    for pattern, availability in NAME_AVAILABILITY_PATTERNS:
        cleaned, replacements = pattern.subn(" ", cleaned)
        if replacements > 0 and detected_availability is None:
            detected_availability = availability
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned, detected_availability


def normalize_availability(value: object, name_availability: str | None = None) -> str | None:
    if name_availability:
        return name_availability
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "В наличии" if value > 0 else "Нет в наличии"
    text = re.sub(r"\s+", " ", str(value).strip()).casefold()
    if not text:
        return None
    if text in {"-", "нет"} or "нет в наличии" in text or "отсут" in text:
        return "Нет в наличии"
    if "под заказ" in text or "заказ" in text:
        return "Под заказ"
    if text in {"+", "да", "есть"} or "в наличии" in text or "налич" in text or "available" in text:
        return "В наличии"
    return None


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value).lower()
    text = text.replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def similarity(left: str, right: str) -> float:
    left_normalized = normalize_text(left)
    right_normalized = normalize_text(right)
    if not left_normalized or not right_normalized:
        return 0.0
    if left_normalized in right_normalized or right_normalized in left_normalized:
        return 1.0
    direct_ratio = SequenceMatcher(None, left_normalized, right_normalized).ratio()
    left_tokens = " ".join(sorted(left_normalized.split()))
    right_tokens = " ".join(sorted(right_normalized.split()))
    token_ratio = SequenceMatcher(None, left_tokens, right_tokens).ratio()
    return max(direct_ratio, token_ratio)


def names_match(expected: str, actual: str, threshold: float = 0.85) -> bool:
    return similarity(expected, actual) >= threshold


def image_code_from_path(path: Path) -> str:
    return path.stem.strip()


def code_lookup_keys(code: str) -> tuple[str, ...]:
    text = normalize_code(code)
    keys = [text]
    if text.isdigit():
        stripped = text.lstrip("0") or "0"
        if stripped not in keys:
            keys.append(stripped)
    return tuple(keys)


def catalog_item_by_code(catalog: dict[str, CatalogItem], code: str) -> CatalogItem | None:
    for key in code_lookup_keys(code):
        item = catalog.get(key)
        if item is not None:
            return item
    if str(code).strip().isdigit():
        target = str(code).strip().lstrip("0") or "0"
        for catalog_code, item in catalog.items():
            if catalog_code.isdigit() and (catalog_code.lstrip("0") or "0") == target:
                return item
    return None


def is_supported_image(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in SUPPORTED_IMAGE_EXTENSIONS


def scan_images(folder: str | Path, catalog: dict[str, CatalogItem] | None = None) -> list[ImageItem]:
    folder_path = Path(folder)
    items: list[ImageItem] = []
    for path in sorted(folder_path.iterdir(), key=lambda item: item.name.lower()):
        if not is_supported_image(path):
            continue
        code = image_code_from_path(path)
        catalog_item = catalog_item_by_code(catalog, code) if catalog else None
        items.append(
            ImageItem(
                code=code,
                path=path,
                catalog_name=catalog_item.name if catalog_item else None,
                price=catalog_item.price if catalog_item else None,
                wholesale_price=catalog_item.wholesale_price if catalog_item else None,
                sko_price=catalog_item.sko_price if catalog_item else None,
                name_ru=None,
                name_kz=None,
                availability=catalog_item.availability if catalog_item else None,
            )
        )
    return items


def processing_items(
    folder: str | Path,
    catalog: dict[str, CatalogItem],
    selected_codes: set[str] | None = None,
) -> list[ImageItem]:
    images_by_code: dict[str, Path] = {}
    for item in scan_images(folder, catalog):
        for key in code_lookup_keys(item.code):
            images_by_code.setdefault(key, item.path)
    items: list[ImageItem] = []
    for code, catalog_item in sorted(catalog.items(), key=lambda item: item[0].lower()):
        if selected_codes is not None and code not in selected_codes:
            continue
        image_path = None
        for key in code_lookup_keys(code):
            image_path = images_by_code.get(key)
            if image_path is not None:
                break
        items.append(
            ImageItem(
                code=code,
                path=image_path,
                catalog_name=catalog_item.name,
                price=catalog_item.price,
                wholesale_price=catalog_item.wholesale_price,
                sko_price=catalog_item.sko_price,
                name_ru=None,
                name_kz=None,
                availability=catalog_item.availability,
            )
        )
    return items


def _header_key(value: object) -> str:
    text = "" if value is None else str(value).lower().replace("ё", "е").replace("_", " ")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def _find_column(headers: Iterable[object], candidates: set[str]) -> int | None:
    normalized_candidates = {_header_key(candidate) for candidate in candidates}
    for index, header in enumerate(headers, start=1):
        key = _header_key(header)
        if key in normalized_candidates:
            return index
    return None


def load_catalog(excel_path: str | Path) -> dict[str, CatalogItem]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise ValueError("В Excel-файле не найдена строка заголовков.")

        code_column = _find_column(header_row, CODE_HEADERS)
        name_column = _find_column(header_row, NAME_HEADERS)
        price_column = _find_column(header_row, PRICE_HEADERS)
        wholesale_price_column = _find_column(header_row, WHOLESALE_PRICE_HEADERS)
        sko_price_column = _find_column(header_row, SKO_PRICE_HEADERS)
        availability_column = _find_column(header_row, AVAILABILITY_HEADERS)
        if code_column is None or name_column is None or price_column is None:
            raise ValueError(
                "Не удалось определить колонки артикула, названия и цены. "
                "Нужны заголовки вроде 'артикул'/'код товара', 'название'/'наименование' и 'Цена'."
            )

        catalog: dict[str, CatalogItem] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            code = normalize_code(row[code_column - 1] if len(row) >= code_column else None)
            raw_name = row[name_column - 1] if len(row) >= name_column else None
            name, name_availability = name_without_availability_marker(raw_name)
            price = normalize_price(row[price_column - 1] if len(row) >= price_column else None)
            wholesale_price = normalize_price(
                row[wholesale_price_column - 1] if wholesale_price_column is not None and len(row) >= wholesale_price_column else None
            )
            sko_price = normalize_price(
                row[sko_price_column - 1] if sko_price_column is not None and len(row) >= sko_price_column else None
            )
            availability_value = row[availability_column - 1] if availability_column is not None and len(row) >= availability_column else None
            availability = normalize_availability(availability_value, name_availability)
            if code and name and code not in catalog:
                catalog[code] = CatalogItem(
                    code=code,
                    name=name,
                    price=price,
                    wholesale_price=wholesale_price,
                    sko_price=sko_price,
                    availability=availability,
                )
        return catalog
    finally:
        workbook.close()


def load_name_catalog(excel_path: str | Path) -> dict[str, NameCatalogItem]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise ValueError("В Excel-файле названий не найдена строка заголовков.")

        code_column = _find_column(header_row, CODE_HEADERS)
        name_ru_column = _find_column(header_row, NAME_RU_HEADERS)
        name_kz_column = _find_column(header_row, NAME_KZ_HEADERS)
        if code_column is None or (name_ru_column is None and name_kz_column is None):
            raise ValueError(
                "Не удалось определить колонки артикула и названий RU/KZ. "
                "Нужны заголовки вроде 'артикул'/'код товара', 'наименование_ru' и/или 'наименование_kz'."
            )

        catalog: dict[str, NameCatalogItem] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            code = normalize_code(row[code_column - 1] if len(row) >= code_column else None)
            if not code or code in catalog:
                continue
            name_ru = normalize_name_cell(row[name_ru_column - 1] if name_ru_column is not None and len(row) >= name_ru_column else None)
            name_kz = normalize_name_cell(row[name_kz_column - 1] if name_kz_column is not None and len(row) >= name_kz_column else None)
            if name_ru or name_kz:
                catalog[code] = NameCatalogItem(code=code, name_ru=name_ru, name_kz=name_kz)
        return catalog
    finally:
        workbook.close()


def normalize_name_cell(value: object) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).strip())
    return text or None
